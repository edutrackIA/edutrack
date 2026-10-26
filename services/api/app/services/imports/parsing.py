import csv
import io
import zipfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

SUPPORTED_EXTENSIONS = {".csv", ".xlsx"}

Row = tuple[int, dict[str, str]]


class UnsupportedFileError(Exception):
    pass


class FileFormatError(Exception):
    pass


def parse_file(filename: str, content: bytes) -> tuple[list[str], list[Row]]:
    """Return normalized headers and data rows numbered as they appear in the file."""
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileError(f"Unsupported file type {extension or '(none)'}")

    table = _read_csv(content) if extension == ".csv" else _read_xlsx(content)
    if not table:
        raise FileFormatError("El archivo está vacío")

    headers = [_normalize_header(value) for value in table[0]]
    rows: list[Row] = []
    for line_number, values in enumerate(table[1:], start=2):
        if not any(value.strip() for value in values):
            continue
        row = {
            header: values[index].strip() if index < len(values) else ""
            for index, header in enumerate(headers)
            if header
        }
        rows.append((line_number, row))
    return headers, rows


def _normalize_header(value: str) -> str:
    return value.strip().lower().replace(" ", "_")


def _read_csv(content: bytes) -> list[list[str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Excel on Spanish Windows exports CSV as cp1252 with ";" separators.
        text = content.decode("cp1252")

    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;")
    except csv.Error:
        dialect = csv.excel
    return list(csv.reader(io.StringIO(text), dialect))


def _read_xlsx(content: bytes) -> list[list[str]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError) as exc:
        raise FileFormatError("No se pudo leer el archivo de Excel") from exc

    try:
        sheet = workbook.active
        if sheet is None:
            raise FileFormatError("El archivo de Excel no tiene hojas")
        return [[_cell_to_text(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _cell_to_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)
